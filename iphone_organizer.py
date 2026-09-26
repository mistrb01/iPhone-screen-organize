"""Organize the iPhone home screen through an Excel round-trip.

Export to Excel - reads the phone's layout and installed apps, looks up App Store categories,
                  caches icons, and writes a workbook. Edit the Target columns in Excel.
Preview         - builds the layout from the workbook and shows it without touching the phone.
Push to iPhone  - backs up the current layout, shows the preview, and writes it on confirm.
Restore Backup  - writes a saved layout back to the phone.

Headless (no window), defaults from iphone_organizer.ini:
    python iphone_organizer.py --headless export [--no-icons]
    python iphone_organizer.py --headless preview [--workbook PATH]
    python iphone_organizer.py --headless push [--workbook PATH] [--yes]
    python iphone_organizer.py --headless restore [--backup PATH] [--yes]
"""

import argparse
import asyncio
import configparser
import os
import json
import plistlib
import re
import signal
import subprocess
import sys
import threading
import tkinter as tk
import urllib.parse
import urllib.request
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, scrolledtext, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from pymobiledevice3.lockdown import create_using_usbmux
from pymobiledevice3.services.installation_proxy import InstallationProxyService
from pymobiledevice3.services.springboard import SpringBoardServicesService

import layout

HERE = Path(__file__).resolve().parent
INI_PATH = HERE / "iphone_organizer.ini"
SECTION = "iphone_organizer"
PATH_RE = re.compile(r"(/[^\s'\"]+\.(?:xlsx|csv|txt|json|plist|docx|zip))")

COLUMNS = [  # (header, width)
    ("Name", 30), ("Kind", 9), ("Category", 18), ("Current Page", 9), ("Current Folder", 22),
    ("Current Order", 9), ("Suggested Folder", 22), ("Target Page", 9), ("Target Folder", 22),
    ("Target Order", 9), ("Bundle ID", 40), ("Key", 40),
]
TARGET_COLS = {"Target Page", "Target Folder", "Target Order"}

HOW_TO = [
    "How to use this workbook",
    "",
    "Edit only the three yellow Target columns on the Apps sheet, save, then click Preview or Push in the tool.",
    "",
    "Target Page: Dock, or a page number (1, 2, 3...). Leave blank (or type Library) to remove the app",
    "    from the home screen. It stays installed and is still in the App Library.",
    "Target Folder: the folder name. Leave blank for a loose icon. Every app with the same folder name",
    "    ends up in one folder, even if they are on different pages now.",
    "Target Order: position on the page, smallest first. Decimals are fine: 3.5 goes between 3 and 4.",
    "    A folder sits at the position of its lowest-numbered app, and apps inside a folder follow",
    "    their Order values. Blank means the end of the page.",
    "",
    "Pages hold 24 icons and the dock holds 4. A page with more than 24 spills onto a new page after it.",
    "Folders show 9 apps per folder page.",
    "Suggested Folder is based on the App Store category. Copy it into Target Folder where it fits.",
    "Apps in rows with a blank Current Page are only in the App Library today.",
    "Widgets can move between pages but can't go into folders or the dock.",
    "Rows sorted or filtered in Excel are fine. Don't delete rows or edit the Key column.",
]


def ts() -> str:
    return datetime.now().strftime("%Y%m%d-%H%M%S")


def to_jsonable(obj):
    if isinstance(obj, bytes):
        return f"<{len(obj)} bytes>"
    if isinstance(obj, datetime):
        return obj.isoformat()
    return str(obj)


def suggested_folder(cat: dict | None, bundle_id: str) -> str:
    if not cat:
        return "Apple" if bundle_id.startswith("com.apple.") else ""
    genres = [g for g in cat.get("genres", []) if g != "Games"]
    if cat.get("genre") == "Games":
        return f"{genres[0]} Games" if genres else "Games"
    return cat.get("genre", "")


class App:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.stop_flag = threading.Event()
        self.busy = False
        self.icon_cache: dict[str, tk.PhotoImage] = {}
        self.output_var = tk.StringVar()
        self.workbook_var = tk.StringVar()
        self.backup_var = tk.StringVar()
        self.icons_var = tk.BooleanVar(value=True)
        self._load_config(quiet=True)
        self.log_path = self._dir("logs") / f"{ts()}-iphone_organizer.txt"
        self.log_file = open(self.log_path, "a", encoding="utf-8")

        root.title("iPhone Home Screen Organizer")
        root.geometry("1050x650")

        fields = ttk.Frame(root, padding=8)
        fields.pack(fill="x")
        ttk.Label(fields, text="Output folder:").grid(row=0, column=0, sticky="w")
        ttk.Entry(fields, textvariable=self.output_var, width=100).grid(row=0, column=1, sticky="ew", padx=4)
        ttk.Button(fields, text="Browse", command=self._browse_output).grid(row=0, column=2)
        for r, (label, var, types, sub) in enumerate([
            ("Workbook:", self.workbook_var, [("Excel", "*.xlsx")], "excel"),
            ("Backup to restore:", self.backup_var, [("plist", "*.plist")], "backups"),
        ], start=1):
            ttk.Label(fields, text=label).grid(row=r, column=0, sticky="w")
            ttk.Entry(fields, textvariable=var, width=100).grid(row=r, column=1, sticky="ew", padx=4)
            ttk.Button(fields, text="Browse", command=lambda v=var, t=types, d=sub: self._browse(v, t, d)).grid(row=r, column=2)
        ttk.Checkbutton(fields, text="Fetch app icons during export (for the preview)", variable=self.icons_var).grid(
            row=3, column=1, sticky="w")
        fields.columnconfigure(1, weight=1)

        self.log_w = scrolledtext.ScrolledText(root, wrap="word", font=("Menlo", 12))
        self.log_w.pack(fill="both", expand=True, padx=8)
        self.log_w.tag_config("link", foreground="blue", underline=True)
        self.log_w.tag_bind("link", "<Button-1>", self._open_link)
        self.log_w.tag_bind("link", "<Enter>", lambda e: self.log_w.config(cursor="hand2"))
        self.log_w.tag_bind("link", "<Leave>", lambda e: self.log_w.config(cursor=""))

        bar = ttk.Frame(root, padding=8)
        bar.pack(fill="x")
        for text, cmd in [
            ("Export to Excel", lambda: self._run(self._export, self.icons_var.get())),
            ("Preview", lambda: self._run(self._prepare, Path(self.workbook_var.get()), False)),
            ("Push to iPhone", lambda: self._run(self._prepare, Path(self.workbook_var.get()), True)),
            ("Restore Backup", lambda: self._run(self._restore, Path(self.backup_var.get()))),
            ("Stop", self._stop),
            ("Save Config", self._save_config),
            ("Load Config", self._load_config),
            ("Copy Log", self._copy_log),
            ("Clear Log", lambda: self.log_w.delete("1.0", "end")),
        ]:
            ttk.Button(bar, text=text, command=cmd).pack(side="left", padx=2)

        if not self.workbook_var.get():
            newest = sorted(self._dir("excel").glob("*-iphone_apps*.xlsx"))
            if newest:
                self.workbook_var.set(str(newest[-1]))
        self.log(f"Log file: {self.log_path}")
        self.log(f"Output folder: {self._dir('')}")
        self.log("Plug in the iPhone and unlock it. Start with Export to Excel.")

    # ---------- logging ----------
    def log(self, msg: str):
        line = f"{ts()}  {msg}"
        try:
            self.log_file.write(line + "\n")
            self.log_file.flush()
            self.root.after(0, self._append, line)
        except (ValueError, RuntimeError, tk.TclError):
            pass  # window already closed

    def _append(self, line: str):
        start = self.log_w.index("end-1c")
        self.log_w.insert("end", line + "\n")
        for m in PATH_RE.finditer(line):
            if Path(m.group(1)).exists():
                self.log_w.tag_add("link", f"{start}+{m.start()}c", f"{start}+{m.end()}c")
        self.log_w.see("end")

    def _open_link(self, event):
        idx = self.log_w.index(f"@{event.x},{event.y}")
        rng = self.log_w.tag_prevrange("link", f"{idx}+1c")
        if not rng:
            return
        path = self.log_w.get(*rng)
        if sys.platform == "darwin":
            subprocess.run(["open", "-R", path])
        elif sys.platform == "win32":
            subprocess.run(["explorer", f"/select,{path}"])

    def _copy_log(self):
        self.root.clipboard_clear()
        self.root.clipboard_append(self.log_w.get("1.0", "end-1c"))

    # ---------- config ----------
    def _save_config(self):
        cfg = configparser.ConfigParser()
        cfg.read(INI_PATH)
        cfg[SECTION] = {
            "output_dir": self.output_var.get(),
            "workbook": self.workbook_var.get(),
            "backup": self.backup_var.get(),
            "fetch_icons": str(self.icons_var.get()),
        }
        with open(INI_PATH, "w") as f:
            cfg.write(f)
        self.log(f"Config saved: {INI_PATH}")

    def _load_config(self, quiet=False):
        cfg = configparser.ConfigParser()
        cfg.read(INI_PATH)
        if cfg.has_section(SECTION):
            s = cfg[SECTION]
            self.output_var.set(s.get("output_dir", ""))
            self.workbook_var.set(s.get("workbook", ""))
            self.backup_var.set(s.get("backup", ""))
            self.icons_var.set(s.getboolean("fetch_icons", True))
        if not quiet:
            self.log("Config loaded.")

    def _dir(self, name: str) -> Path:
        """A subfolder of the output folder (blank = next to this script), created if missing."""
        d = Path(self.output_var.get().strip() or HERE).expanduser() / name
        d.mkdir(parents=True, exist_ok=True)
        return d

    def _browse(self, var, types, sub):
        p = filedialog.askopenfilename(initialdir=self._dir(sub), filetypes=types)
        if p:
            var.set(p)

    def _browse_output(self):
        p = filedialog.askdirectory(initialdir=self._dir(""), mustexist=False)
        if p:
            self.output_var.set(p)
            self.log(f"Output folder: {p}. Click Save Config to keep it. This session's log stays at {self.log_path}")

    # ---------- run plumbing ----------
    def _stop(self):
        self.stop_flag.set()
        self.log("Stop requested. The current item will finish, then the operation ends.")

    def _run(self, coro_fn, *args):
        if self.busy:
            self.log("Already running.")
            return
        self.busy = True
        self.stop_flag.clear()

        def worker():
            try:
                asyncio.run(coro_fn(*args))
            except Exception as e:
                self.log(f"ERROR: {type(e).__name__}: {e}")
            finally:
                self.busy = False

        threading.Thread(target=worker, daemon=True).start()

    def _ui(self, fn, *args):
        self.root.after(0, fn, *args)

    def _warn(self, title: str, msg: str):
        self.root.after(0, lambda: messagebox.showwarning(title, msg))

    def _ask(self, title: str, msg: str) -> bool:
        done, answer = threading.Event(), []

        def show():
            answer.append(messagebox.askokcancel(title, msg))
            done.set()

        self.root.after(0, show)
        done.wait()
        return answer[0]

    async def _connect(self):
        self.log("Connecting over USB...")
        lockdown = await create_using_usbmux()
        info = lockdown.short_info
        self.log(f"Connected: {info.get('DeviceName')} ({info.get('ProductType')}), iOS {info.get('ProductVersion')}")
        return lockdown, SpringBoardServicesService(lockdown)

    def _backup(self, state: list, tag: str) -> Path:
        stamp = ts()
        plist_path = self._dir("backups") / f"{stamp}-icon_state-{tag}.plist"
        plist_path.write_bytes(plistlib.dumps(state, fmt=plistlib.FMT_BINARY))
        (self._dir("json") / f"{stamp}-icon_state-{tag}.json").write_text(
            json.dumps(state, indent=2, default=to_jsonable), encoding="utf-8")
        self.log(f"Layout backup: {plist_path}")
        return plist_path

    # ---------- export ----------
    async def _export(self, fetch_icons: bool):
        lockdown, sb = await self._connect()
        rows, cats = [], {}
        try:
            state = await sb.get_icon_state()
            self._ui(self.backup_var.set, str(self._backup(state, "export")))
            rows = layout.flatten(state)
            on_home = {r["key"] for r in rows}
            apps = await InstallationProxyService(lockdown).get_apps(application_type="User")
            for bid, info in sorted(apps.items(), key=lambda kv: kv[1].get("CFBundleDisplayName") or kv[0]):
                if bid not in on_home:
                    rows.append({"key": bid, "bundle_id": bid, "kind": "App", "page": None, "folder": None,
                                 "order": None, "name": info.get("CFBundleDisplayName") or info.get("CFBundleName") or bid})
            self.log(f"{len(on_home)} items on the home screen, {len(rows) - len(on_home)} apps only in the App Library.")

            cats = self._categories({r["bundle_id"] for r in rows if r["kind"] == "App"})
            if fetch_icons and not self.stop_flag.is_set():
                await self._fetch_icons(sb, sorted({r["bundle_id"] for r in rows}))
        except Exception as e:
            self.log(f"ERROR during export: {type(e).__name__}: {e}")
            if rows:
                self.log("Writing a partial workbook with what was collected.")
                path = self._write_workbook(rows, cats)
                self._warn("Partial export", f"Partial workbook saved:\n{path}")
            return
        finally:
            await lockdown.close()
        path = self._write_workbook(rows, cats)
        self._ui(self.workbook_var.set, str(path))
        self.log("Edit the yellow Target columns, save, then click Preview.")

    def _categories(self, bundle_ids: set[str]) -> dict:
        cached = sorted(self._dir("json").glob("*-app_categories.json"))
        cats = json.loads(cached[-1].read_text()) if cached else {}
        todo = sorted(b for b in bundle_ids if b not in cats and not b.startswith("com.apple."))
        self.log(f"App Store categories: looking up {len(todo)} (Apple apps and cached ones skipped)...")
        for i in range(0, len(todo), 50):
            if self.stop_flag.is_set():
                self.log("Stopped during category lookup.")
                break
            batch = todo[i:i + 50]
            url = "https://itunes.apple.com/lookup?country=us&bundleId=" + urllib.parse.quote(",".join(batch))
            try:
                with urllib.request.urlopen(url, timeout=30) as resp:
                    data = json.load(resp)
            except Exception as e:
                self.log(f"Category lookup failed for a batch: {e}")
                continue
            for r in data.get("results", []):
                cats[r["bundleId"]] = {"genre": r.get("primaryGenreName"), "genres": r.get("genres", [])}
            for b in batch:
                cats.setdefault(b, {})
        path = self._dir("json") / f"{ts()}-app_categories.json"
        path.write_text(json.dumps(cats, indent=2), encoding="utf-8")
        self.log(f"Categories saved: {path}")
        return cats

    async def _fetch_icons(self, sb, bundle_ids: list[str]):
        icon_dir = self._dir("icons")
        todo = [b for b in bundle_ids if not (icon_dir / f"{b}.png").exists()]
        self.log(f"Icons: {len(bundle_ids) - len(todo)} cached, fetching {len(todo)}...")
        for n, bid in enumerate(todo, start=1):
            if self.stop_flag.is_set():
                self.log(f"Stopped after {n - 1} icons. The rest can be fetched on the next export.")
                return
            try:
                png = await sb.get_icon_pngdata(bid)
                if png:
                    (icon_dir / f"{bid}.png").write_bytes(png)
            except Exception as e:
                self.log(f"Icon failed for {bid}: {e}")
            if n % 50 == 0:
                self.log(f"  {n}/{len(todo)} icons")
        self.log("Icons done.")

    def _write_workbook(self, rows: list[dict], cats: dict) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "Apps"
        ws.append([h for h, _ in COLUMNS])

        def sort_key(r):
            p = r["page"]
            return (2, 0, r["name"].lower()) if p is None else (0 if p == "Dock" else 1, 0 if p == "Dock" else p, r["order"])

        for r in sorted(rows, key=sort_key):
            cat = cats.get(r["bundle_id"])
            ws.append([
                r["name"], r["kind"],
                (cat or {}).get("genre") or ("Apple" if r["bundle_id"].startswith("com.apple.") else ""),
                r["page"], r["folder"], r["order"],
                "" if r["kind"] == "Widget" else suggested_folder(cat, r["bundle_id"]),
                r["page"], r["folder"], r["order"],
                r["bundle_id"], r["key"],
            ])

        big = Font(size=14)
        bold = Font(size=14, bold=True)
        yellow = PatternFill("solid", fgColor="FFF2CC")
        target_idx = {i for i, (h, _) in enumerate(COLUMNS, start=1) if h in TARGET_COLS}
        for row in ws.iter_rows():
            for c in row:
                c.font = bold if c.row == 1 else big
                if c.row > 1 and c.column in target_idx:
                    c.fill = yellow
        for i, (_, w) in enumerate(COLUMNS, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w * 1.25
        ref = f"A1:{ws.cell(ws.max_row, len(COLUMNS)).coordinate}"
        table = Table(displayName="Apps", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        ws.add_table(table)
        ws.freeze_panes = "B2"

        how = wb.create_sheet("How To")
        for line in HOW_TO:
            how.append([line])
        for row in how.iter_rows():
            for c in row:
                c.font = bold if c.row == 1 else big
        how.column_dimensions["A"].width = 130

        path = self._dir("excel") / f"{ts()}-iphone_apps.xlsx"
        wb.save(path)
        self.log(f"Workbook: {path}")
        return path

    # ---------- read workbook ----------
    def _read_workbook(self, path: Path) -> list[dict] | None:
        if not path.exists():
            self.log(f"Workbook not found: {path}")
            return None
        if (path.parent / f"~${path.name}").exists():
            self.log("The workbook is open in Excel. Save and close it first, so the tool reads your latest edits.")
            return None
        ws = load_workbook(path, data_only=True)["Apps"]
        it = ws.iter_rows(values_only=True)
        idx = {h: i for i, h in enumerate(next(it))}
        rows, bad = [], []
        for n, v in enumerate(it, start=2):
            if not v or not v[idx["Key"]]:
                continue
            page = v[idx["Target Page"]]
            if page is None or str(page).strip().lower() in ("", "library"):
                page = None
            elif str(page).strip().lower() == "dock":
                page = "Dock"
            else:
                try:
                    page = int(float(page))
                    if page < 1:
                        raise ValueError
                except ValueError:
                    bad.append(f"Row {n} ({v[idx['Name']]}): Target Page '{page}' is not Dock, a page number, or blank.")
                    continue
            order = v[idx["Target Order"]]
            try:
                order = float(order) if order not in (None, "") else None
            except ValueError:
                bad.append(f"Row {n} ({v[idx['Name']]}): Target Order '{order}' is not a number.")
                continue
            folder = v[idx["Target Folder"]]
            rows.append({
                "key": str(v[idx["Key"]]),
                "bundle_id": str(v[idx["Bundle ID"]]),
                "name": str(v[idx["Name"]]),
                "kind": v[idx["Kind"]],
                "page": page,
                "folder": str(folder).strip() if folder not in (None, "") else None,
                "order": order,
            })
        for b in bad:
            self.log(b)
        if bad:
            self.log("Fix those rows and try again. Nothing was sent to the phone.")
            return None
        self.log(f"Read {len(rows)} rows from {path}")
        return rows

    # ---------- preview / push ----------
    async def _prepare(self, workbook: Path, push: bool):
        rows = self._read_workbook(workbook)
        if rows is None:
            return
        lockdown, sb = await self._connect()
        try:
            before = await sb.get_icon_state()
        finally:
            await lockdown.close()
        new_state, hidden, warnings, errors = layout.build(rows, before)
        for w in warnings:
            self.log(f"WARNING: {w}")
        for e in errors:
            self.log(f"ERROR: {e}")
        if errors:
            self.log("Fix the errors above and try again. Nothing was sent to the phone.")
            return
        summary = layout.summarize_changes(before, new_state)
        for line in summary:
            self.log(line)
        if layout.signature(new_state) == layout.signature(before):
            self.log("The workbook matches the phone's current layout. Nothing to change.")
            if push:
                return
        backup = None
        if push:
            backup = self._backup(before, "before-push")
            self._ui(self.backup_var.set, str(backup))
        self._ui(self._show_preview, before, new_state, hidden, summary + warnings, backup)

    def _icon(self, bundle_id: str | None) -> tk.PhotoImage | None:
        if not bundle_id:
            return None
        if bundle_id not in self.icon_cache:
            f = self._dir("icons") / f"{bundle_id}.png"
            try:
                img = tk.PhotoImage(file=str(f)) if f.exists() else None
                self.icon_cache[bundle_id] = img.subsample(max(1, img.width() // 48)) if img else None
            except tk.TclError:
                self.icon_cache[bundle_id] = None
        return self.icon_cache[bundle_id]

    def _show_preview(self, before, new_state, hidden, notes, backup):
        old = {r["key"]: (r["page"], r["folder"]) for r in layout.flatten(before)}
        changed = {r["key"] for r in layout.flatten(new_state) if old.get(r["key"]) != (r["page"], r["folder"])}

        win = tk.Toplevel(self.root)
        win.title("Push preview" if backup else "Preview")
        win.geometry("1400x900")
        ttk.Label(win, text="\n".join(notes) + "\nHighlighted icons change page or folder. Click a folder to list its apps.",
                  justify="left", padding=8).pack(fill="x")

        btns = ttk.Frame(win, padding=8)
        btns.pack(side="bottom", fill="x")
        if backup:
            def do_push():
                if messagebox.askokcancel("Push to iPhone", f"Write this layout to the phone?\n\nUndo with Restore Backup:\n{backup}", parent=win):
                    win.destroy()
                    self._run(self._push, new_state, hidden)
            ttk.Button(btns, text="Push to iPhone", command=do_push).pack(side="left", padx=4)
        ttk.Button(btns, text="Close", command=win.destroy).pack(side="left", padx=4)

        canvas = tk.Canvas(win, highlightthickness=0, yscrollincrement=1)  # 1 unit = 1 pixel
        sbar = ttk.Scrollbar(win, orient="vertical", command=canvas.yview)
        inner = ttk.Frame(canvas)
        inner.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=inner, anchor="nw")
        canvas.configure(yscrollcommand=sbar.set)
        sbar.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)

        # Tk 9: a wheel notch is delta 120 everywhere; macOS trackpads send <TouchpadScroll> instead.
        win.bind("<MouseWheel>", lambda e: canvas.yview_scroll(int(-e.delta * 40 / 120) or (-1 if e.delta > 0 else 1), "units"))

        def touchpad(e):
            _, dy = (float(v) for v in win.tk.splitlist(win.tk.call("tk::PreciseScrollDeltas", e.delta)))
            canvas.yview_scroll(int(-dy * 2), "units")

        win.bind("<TouchpadScroll>", touchpad)

        per_row = 4
        for pi, page in enumerate(new_state):
            title = "Dock" if pi == 0 else f"Page {pi}  ({len(page)} items)"
            if pi > layout.MAX_PAGES:
                title += "  OVER THE 15-PAGE LIMIT: App Library only"
            lf = ttk.LabelFrame(inner, text=title, padding=4)
            slot = pi - 1 if pi else 0
            lf.grid(row=0 if pi == 0 else 1 + slot // per_row, column=0 if pi == 0 else slot % per_row,
                    columnspan=per_row if pi == 0 else 1, sticky="nw", padx=6, pady=6)
            for n, it in enumerate(page):
                self._tile(lf, it, changed).grid(row=n // 4, column=n % 4, padx=2, pady=2, sticky="n")
        if hidden:
            rows = 2 + (len(new_state) - 2) // per_row
            lf = ttk.LabelFrame(inner, text=f"App Library only ({len(hidden)})", padding=4)
            lf.grid(row=rows, column=0, columnspan=per_row, sticky="nw", padx=6, pady=6)
            ttk.Label(lf, text=", ".join(sorted(hidden, key=str.lower)), wraplength=1300, justify="left").pack()

    def _tile(self, parent, it: dict, changed: set) -> tk.Frame:
        f = tk.Frame(parent, width=80, height=78)
        f.grid_propagate(False)
        f.pack_propagate(False)
        if layout.is_folder(it):
            members = [m for sub in it["iconLists"] for m in sub]
            hot = any(layout.item_key(m) in changed for m in members)
            lbl = tk.Label(f, text=f"{it['displayName']}\n({len(members)})", bg="#ffe08a" if hot else "#d8dde6", fg="black",
                           wraplength=74, font=("Helvetica", 10, "bold"))
            lbl.pack(fill="both", expand=True)
            names = "\n".join(m.get("displayName", "?") for m in members)
            lbl.bind("<Button-1>", lambda e: messagebox.showinfo(it["displayName"], names, parent=parent.winfo_toplevel()))
            return f
        hot = layout.item_key(it) in changed
        if hot:
            f.configure(bg="#ffe08a")
        img = self._icon(it.get("bundleIdentifier"))
        name = it.get("displayName") or "Widget"
        if img:
            tk.Label(f, image=img, bg=f["bg"]).pack()
        else:
            tk.Label(f, text="[widget]" if layout.kind_of(it) == "Widget" else "[ ]", height=2, bg=f["bg"],
                     **({"fg": "black"} if hot else {})).pack()
        tk.Label(f, text=name[:22], wraplength=76, font=("Helvetica", 9), bg=f["bg"], **({"fg": "black"} if hot else {})).pack()
        return f

    async def _push(self, new_state: list, hidden: list[str]):
        lockdown, sb = await self._connect()
        try:
            self.log("Writing the new layout...")
            await sb.set_icon_state(new_state)
            await asyncio.sleep(3)
            after = await sb.get_icon_state()
        finally:
            await lockdown.close()
        self._backup(after, "after-push")
        if layout.signature(after) == layout.signature(new_state):
            self.log("RESULT: the phone now matches the preview.")
            return
        self.log("RESULT: the phone's layout differs from the preview. iOS adjusted it:")
        want = {r["key"]: r for r in layout.flatten(new_state)}
        got = {r["key"]: r for r in layout.flatten(after)}
        diffs = 0
        for k, r in got.items():
            w = want.get(k)
            if w is None:
                where = "App Library" if r["name"] in hidden else "unplanned"
                self.log(f"  {r['name']}: expected {where}, is on page {r['page']}")
                diffs += 1
            elif (w["page"], w["folder"]) != (r["page"], r["folder"]):
                self.log(f"  {r['name']}: expected page {w['page']} / {w['folder'] or 'no folder'}, "
                         f"got page {r['page']} / {r['folder'] or 'no folder'}")
                diffs += 1
        for k, w in want.items():
            if k not in got:
                self.log(f"  {w['name']}: expected page {w['page']}, missing from the home screen")
                diffs += 1
        if not diffs:
            self.log("  Every app is on the right page and in the right folder. Only the order within pages differs.")
        self.log("Use Restore Backup to go back to the previous layout if needed.")

    async def _restore(self, path: Path):
        if not path.exists():
            self.log(f"Backup not found: {path}")
            return
        state = plistlib.loads(path.read_bytes())
        if not self._ask("Restore Backup", f"Write this saved layout to the phone?\n\n{path}"):
            return
        lockdown, sb = await self._connect()
        try:
            self.log(f"Restoring {path} ...")
            await sb.set_icon_state(state)
            await asyncio.sleep(3)
            after = await sb.get_icon_state()
        finally:
            await lockdown.close()
        ok = layout.signature(after) == layout.signature(state)
        self.log("Restore verified." if ok else "Restore sent, but the re-read layout differs. Check the phone.")

    def close(self):
        self.stop_flag.set()
        self.log_file.close()
        self.root.destroy()


class Var:
    def __init__(self, value=None):
        self.value = value

    def get(self):
        return self.value

    def set(self, value):
        self.value = value


class Headless(App):
    """Same operations as the GUI, driven from the command line."""

    def __init__(self, assume_yes: bool):
        self.stop_flag = threading.Event()
        self.busy = False
        self.assume_yes = assume_yes
        self.pending_push = None
        self.output_var, self.workbook_var, self.backup_var, self.icons_var = Var(""), Var(""), Var(""), Var(True)
        self._load_config(quiet=True)
        self.log_path = self._dir("logs") / f"{ts()}-iphone_organizer.txt"
        self.log_file = open(self.log_path, "a", encoding="utf-8")
        self.log(f"Log file: {self.log_path}")

    def log(self, msg: str):
        line = f"{ts()}  {msg}"
        print(line, flush=True)
        if not self.log_file.closed:
            self.log_file.write(line + "\n")
            self.log_file.flush()

    def _ui(self, fn, *args):
        fn(*args)

    def _warn(self, title: str, msg: str):
        self.log(f"WARNING: {title}: {msg}")

    def _ask(self, title: str, msg: str) -> bool:
        if self.assume_yes:
            return True
        return input(f"{title}: {msg}\nType y to continue: ").strip().lower() == "y"

    def _show_preview(self, before, new_state, hidden, notes, backup):
        old = {r["key"]: (r["page"], r["folder"]) for r in layout.flatten(before)}
        changed = {r["key"] for r in layout.flatten(new_state) if old.get(r["key"]) != (r["page"], r["folder"])}
        print("\nPreview (* = changes page or folder)")
        for pi, page in enumerate(new_state):
            title = "Dock" if pi == 0 else f"Page {pi}"
            if pi > layout.MAX_PAGES:
                title += " (OVER THE 15-PAGE LIMIT: App Library only)"
            print(f"\n{title} ({len(page)} items)")
            for it in page:
                if layout.is_folder(it):
                    members = [m for sub in it["iconLists"] for m in sub]
                    names = ", ".join(("*" if layout.item_key(m) in changed else "") + m.get("displayName", "?") for m in members)
                    print(f"  [{it['displayName']}] ({len(members)}): {names}")
                else:
                    mark = "*" if layout.item_key(it) in changed else " "
                    print(f" {mark}{it.get('displayName') or 'Widget'}")
        if hidden:
            print(f"\nApp Library only ({len(hidden)}): " + ", ".join(sorted(hidden, key=str.lower)))
        print()
        if backup and self._ask("Push to iPhone", f"Write this layout to the phone? Undo with restore --backup {backup}"):
            self.pending_push = (new_state, hidden)
        elif backup:
            self.log("Push cancelled. Nothing was sent to the phone.")

    def close(self):
        self.log_file.close()


def run_headless(args):
    app = Headless(args.yes)
    presses = {"n": 0}

    def on_sigint(*_):
        presses["n"] += 1
        if presses["n"] == 1:
            app._stop()
        elif presses["n"] == 2:
            raise KeyboardInterrupt
        else:
            os._exit(1)

    signal.signal(signal.SIGINT, on_sigint)
    workbook = Path(args.workbook or app.workbook_var.get())
    backup = Path(args.backup or app.backup_var.get())
    try:
        if args.action == "export":
            asyncio.run(app._export(not args.no_icons))
        elif args.action in ("preview", "push"):
            asyncio.run(app._prepare(workbook, args.action == "push"))
            if app.pending_push:
                asyncio.run(app._push(*app.pending_push))
        elif args.action == "restore":
            asyncio.run(app._restore(backup))
        app._save_config()
    except KeyboardInterrupt:
        app.log("Stopped. Files written so far are kept.")
        sys.exit(130)
    except Exception as e:
        app.log(f"ERROR: {type(e).__name__}: {e}")
        sys.exit(1)
    finally:
        app.close()


def main():
    parser = argparse.ArgumentParser(description="Organize the iPhone home screen through an Excel round-trip.")
    parser.add_argument("--headless", action="store_true", help="run one action without the window")
    parser.add_argument("action", nargs="?", choices=["export", "preview", "push", "restore"])
    parser.add_argument("--workbook", help="workbook for preview/push (default: from the .ini)")
    parser.add_argument("--backup", help="backup .plist for restore (default: from the .ini)")
    parser.add_argument("--no-icons", action="store_true", help="skip icon download during export")
    parser.add_argument("--yes", action="store_true", help="push/restore without asking")
    args = parser.parse_args()
    if args.headless:
        if not args.action:
            parser.error("--headless needs an action: export, preview, push or restore")
        run_headless(args)
        return
    root = tk.Tk()
    app = App(root)
    presses = {"n": 0}

    def on_sigint(*_):
        presses["n"] += 1
        if presses["n"] >= 3:
            sys.exit(1)
        app._stop()
        if presses["n"] == 2 or not app.busy:
            root.after(0, app.close)

    signal.signal(signal.SIGINT, on_sigint)
    root.protocol("WM_DELETE_WINDOW", app.close)

    def tick():  # lets Python see Ctrl-C while Tk is idle
        root.after(250, tick)

    tick()
    root.mainloop()


if __name__ == "__main__":
    main()
